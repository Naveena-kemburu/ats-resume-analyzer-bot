import os
import logging
import PyPDF2
import docx

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = ('.pdf', '.docx', '.doc', '.txt')

def is_supported_file(filename: str) -> bool:
    """Check if the filename has a supported document extension"""
    if not filename:
        return False
    return filename.lower().endswith(SUPPORTED_EXTENSIONS)

def extract_text_from_file(file_path: str) -> str:
    """Extract text from PDF, DOCX, or TXT file"""
    if not os.path.exists(file_path):
        logger.error(f"File not found: {file_path}")
        return ""
    
    ext = os.path.splitext(file_path)[1].lower()
    
    try:
        if ext == '.pdf':
            return extract_text_from_pdf(file_path)
        elif ext in ('.docx', '.doc'):
            return extract_text_from_docx(file_path)
        elif ext in ('.txt', '.md'):
            return extract_text_from_txt(file_path)
        else:
            logger.warning(f"Unsupported file extension: {ext}")
            return ""
    except Exception as e:
        logger.error(f"Error extracting text from {file_path}: {str(e)}")
        return ""

def extract_text_from_pdf(pdf_path: str) -> str:
    """Extract text from PDF file"""
    try:
        with open(pdf_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            text = []
            for page in pdf_reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text.append(page_text)
            return "\n".join(text).strip()
    except Exception as e:
        logger.error(f"Error extracting PDF text from {pdf_path}: {str(e)}")
        return ""

def extract_text_from_docx(docx_path: str) -> str:
    """Extract text from DOCX file including paragraphs and tables"""
    try:
        doc = docx.Document(docx_path)
        text = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text.append(paragraph.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text.append(" | ".join(row_text))
        return "\n".join(text).strip()
    except Exception as e:
        logger.error(f"Error extracting DOCX text from {docx_path}: {str(e)}")
        return ""

def extract_text_from_txt(txt_path: str) -> str:
    """Extract text from plain text file"""
    for enc in ('utf-8', 'latin-1', 'cp1252'):
        try:
            with open(txt_path, 'r', encoding=enc) as f:
                return f.read().strip()
        except Exception:
            continue
    return ""
